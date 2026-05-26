# parser.py
import re, uuid
from docx import Document
from models import Dish, Category, MealTime, FullMenu
from datetime import datetime
from typing import Optional


MONTHS_RU = {
    'января': 1, 'февраля': 2, 'марта': 3, 'апреля': 4, 'мая': 5, 'июня': 6,
    'июля': 7, 'августа': 8, 'сентября': 9, 'октября': 10, 'ноября': 11, 'декабря': 12
}


def extract_menu_date(doc) -> Optional[str]:
    """Ищет дату в формате 'Меню на 05 февраля 2026 года'"""
    text_parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            text_parts.extend([cell.text for cell in row.cells])
    
    full_text = " ".join(text_parts)
    pattern = r"Меню\s+на\s+(\d{1,2})\s+(января|февраля|марта|апреля|мая|июня|июля|августа|сентября|октября|ноября|декабря)\s+(\d{4})"
    match = re.search(pattern, full_text, re.IGNORECASE)
    
    if match:
        day, month_str, year = match.groups()
        month = MONTHS_RU.get(month_str.lower())
        if month:
            try:
                dt = datetime(int(year), month, int(day))
                return dt.isoformat()  # Вернёт "2026-02-05T00:00:00"
            except ValueError:
                pass
    return None

def parse_price(price_str: str) -> float:
    """Парсит цены вида '38-00', '10 -00', '120,50' в float"""
    if not price_str or not price_str.strip():
        return 0.0
    
    cleaned = price_str.strip().replace(" ", "")
    # Формат XX-YY или XX,YY
    match = re.match(r'^(\d+)[\-,](\d{2})$', cleaned)
    if match:
        return float(f"{match.group(1)}.{match.group(2)}")
        
    # Фоллбэк: обычное число
    try:
        return float(cleaned.replace(",", "."))
    except ValueError:
        return 0.0

def parse_docx(file_path: str) -> FullMenu:
    doc = Document(file_path)
    menu = FullMenu()
    
    
    # 📅 Извлекаем дату актуальности
    menu.valid_date = extract_menu_date(doc)
    
    # 🔥 КЛЮЧЕВОЕ: первая таблица = завтрак, вторая = обед
    table_index = 0
    
    for table in doc.tables:
        table_index += 1
        current_section = MealTime.breakfast if table_index == 1 else MealTime.lunch
        current_category = Category.unknown
        
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) < 4:
                continue
            
            full_text = " ".join(cells).lower()
            
            # Для ЗАВТРАКА (первая таблица)
            if current_section == MealTime.breakfast:
                # Пропускаем заголовки и мусор
                if any(x in full_text for x in ["завтрак", "меню на", "года", "зав."]):
                    continue
                
                name = cells[1].strip()
                price_str = cells[3] if len(cells) > 3 else cells[-1]
                price = parse_price(price_str)
                
                # Фильтруем мусор
                if price <= 0 or len(name) < 3:
                    continue
                
                dish = Dish(
                    uuid=uuid.uuid4().hex[:10],
                    name=name,
                    price=price,
                    category=Category.unknown  # Завтрак без категорий
                )
                menu.breakfast.dishes.append(dish)
            
            # Для ОБЕДА (вторая таблица)
            elif current_section == MealTime.lunch:
                # Проверяем, не является ли строка заголовком категории
                # Строки вида "| |Салаты|Салаты|Салаты|"
                if len(cells) >= 3 and cells[1] and cells[1] == cells[2]:
                    cat_name = cells[1].lower()
                    for cat in Category:
                        if cat.value.lower() in cat_name:
                            current_category = cat
                            break
                    continue  # Пропускаем строку-категорию
                
                # Парсим блюдо
                name = cells[1].strip()
                price_str = cells[3] if len(cells) > 3 else cells[-1]
                price = parse_price(price_str)
                
                # Фильтруем мусор
                skip_phrases = ["зав.", "производством", "меню", "года"]
                if price <= 0 or len(name) < 3 or any(p in name.lower() for p in skip_phrases):
                    continue
                
                dish = Dish(
                    uuid=uuid.uuid4().hex[:10],
                    name=name,
                    price=price,
                    category=current_category
                )
                menu.lunch.dishes.append(dish)
    
    # Сортируем блюда по категориям (для обеда)
    menu.breakfast.dishes.sort(key=lambda d: d.category.order_index)
    menu.lunch.dishes.sort(key=lambda d: d.category.order_index)
    return menu